import openpyxl
import json
import re
import uuid
import datetime

EXCEL_PATH = r'C:/Users/salma/.gemini/antigravity/brain/4b28ffd2-b82b-4f13-a398-9117f9244957/.user_uploaded/media_1790456866427.xlsx'
OUT_SQL = r'C:/Users/salma/Music/Dhako/dhako-app/supabase/migrations/004_seed_inventory.sql'

def escape_sql(val):
    if val is None:
        return 'NULL'
    if isinstance(val, (int, float)):
        return str(val)
    # escape single quotes
    return "'" + str(val).replace("'", "''") + "'"

def infer_category(name):
    name_lower = name.lower()
    mapping = {
        'filt': 'Filtarada',
        'shaka': 'Shaka Khafis',
        'rabadh': 'Rabadh Kaabaan',
        'buush': 'Buush',
        'kaanweys': 'Kaanweys',
        'boolbeer': 'Boolbeerin',
        'muraad': 'Muraadyad Haad',
        'khashaa': 'Khashaafad',
        'indho': 'Laydhadhka',
        'lambad': 'Laydhadhka',
        'guluub': 'Laydhadhka',
        'lamdad': 'Lamdado',
        'bool': 'Bool / Baanad',
        'baanad': 'Bool / Baanad',
        'saliid': 'Saliid / Dacawo',
        'dawo': 'Saliid / Dacawo',
        'tuubo': 'Tuubo',
        'abwaal': 'Abwaal',
        'xidhiidhiye': 'Xidhiidhiye',
        'dhanbaraas': 'Qalab / Tools',
        'dhagax': 'Qalab / Tools',
        'siiq': 'Siiq',
        'kiliish': 'Kiliish / Okiyo',
        'okiyo': 'Kiliish / Okiyo',
        'connector': 'Koronto / Electrical',
        'iswiij': 'Koronto / Electrical',
        'fiyuus': 'Koronto / Electrical'
    }
    for key, cat in mapping.items():
        if key in name_lower:
            return cat
    return 'Guud / General'

def infer_model(name):
    name_lower = name.lower()
    if 'tx400' in name_lower or 'tx 400' in name_lower:
        return 'TX400'
    if '371' in name_lower:
        return '371'
    if '336' in name_lower:
        return '336'
    if 'sino' in name_lower or 'zino' in name_lower:
        return 'Sino'
    return 'General'

def infer_grade(name):
    name_lower = name.lower()
    if 'org' in name_lower or 'asli' in name_lower:
        return 'original'
    if 'copy' in name_lower or 'koobi' in name_lower:
        return 'copy'
    return 'unknown'

def extract_part_code(name):
    # look for numbers in parens or trailing numbers
    m = re.search(r'\((\d[a-zA-Z0-9]*)\)', name)
    if m:
        return m.group(1)
    # trailing numbers separated by dash
    m = re.search(r'-\s*([A-Z0-9]+)\s*$', name)
    if m:
        return m.group(1)
    return None

def run():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Sheet1']
    
    sql_lines = [
        "-- MIGRATION: SEED INVENTORY FROM EXCEL",
        "DO $$",
        "DECLARE",
        "  cat_id UUID;",
        "  prod_id UUID;",
        "BEGIN"
    ]
    
    products = {} # track by name to merge duplicates
    
    for row_idx in range(2, ws.max_row + 1):
        name = ws.cell(row=row_idx, column=1).value
        qty = ws.cell(row=row_idx, column=2).value
        cost = ws.cell(row=row_idx, column=3).value
        sell = ws.cell(row=row_idx, column=5).value
        
        if not name:
            continue
            
        name = str(name).strip()
        qty = int(qty) if qty is not None else 0
        cost = float(cost) if cost is not None else 0.0
        sell = float(sell) if sell is not None else 0.0
        
        if name in products:
            products[name]['qty'] += qty
        else:
            products[name] = {
                'qty': qty,
                'cost': cost,
                'sell': sell
            }
            
    # Generate SQL for each product
    for name, data in products.items():
        cat = infer_category(name)
        model = infer_model(name)
        grade = infer_grade(name)
        code = extract_part_code(name)
        
        sql_lines.append(f"  -- {name}")
        sql_lines.append(f"  SELECT id INTO cat_id FROM categories WHERE name_so = '{cat}' LIMIT 1;")
        sql_lines.append(f"  IF cat_id IS NULL THEN SELECT id INTO cat_id FROM categories WHERE name_so = 'Guud / General' LIMIT 1; END IF;")
        
        sql_lines.append(f"  INSERT INTO products (name, part_code, category_id, vehicle_model, quality_grade, cost_price, selling_price, quantity_on_hand, min_stock_level, status)")
        sql_lines.append(f"  VALUES ({escape_sql(name)}, {escape_sql(code)}, cat_id, {escape_sql(model)}, {escape_sql(grade)}, {data['cost']}, {data['sell']}, {data['qty']}, 5, 'active')")
        sql_lines.append(f"  RETURNING id INTO prod_id;")
        
        # Insert initial stock movement
        if data['qty'] > 0:
            sql_lines.append(f"  INSERT INTO stock_movements (product_id, movement_type, quantity_change, quantity_before, quantity_after, notes)")
            sql_lines.append(f"  VALUES (prod_id, 'initial_stock', {data['qty']}, 0, {data['qty']}, 'Excel Migration');")
            
    sql_lines.append("END $$;")
    
    with open(OUT_SQL, 'w', encoding='utf-8') as f:
        f.write('\n'.join(sql_lines))
        
    print(f"Generated {OUT_SQL} with {len(products)} unique products.")

if __name__ == '__main__':
    run()
